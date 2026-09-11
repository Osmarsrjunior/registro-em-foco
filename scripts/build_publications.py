"""Gera artigo PDF/DOCX e páginas estáticas a partir dos textos versionados.

Uso: python scripts/build_publications.py --poppler /caminho/pdftoppm
As figuras PNG versionadas permitem repetir a publicação sem Poppler.
"""
import argparse
import csv
import html
import re
import shutil
import subprocess
from pathlib import Path
from reportlab.graphics.shapes import Drawing, Rect, String, Line
from reportlab.graphics import renderPDF
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, Table, TableStyle, KeepTogether
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import reportlab
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
NAVY=colors.HexColor('#112643'); TEAL=colors.HexColor('#007b73')
FONTS=Path(reportlab.__file__).parent/'fonts'
pdfmetrics.registerFont(TTFont('RF',str(FONTS/'Vera.ttf')))
pdfmetrics.registerFont(TTFont('RF-Bold',str(FONTS/'VeraBd.ttf')))


def blocks(text):
    """Parser pequeno, deliberadamente limitado ao Markdown deste projeto."""
    lines=text.splitlines(); i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith('#'):
            n=len(line)-len(line.lstrip('#'));yield ('heading',n,line[n:].strip());i+=1
        elif line.startswith('!['):
            m=re.fullmatch(r'!\[(.*?)\]\((.*?)\)',line);yield ('image',m[1],m[2]);i+=1
        elif line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+',x) for x in cells): rows.append(cells)
                i+=1
            yield ('table',rows)
        else:
            paragraph=[line];i+=1
            while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','![')):
                paragraph.append(lines[i].strip());i+=1
            yield ('paragraph',' '.join(paragraph))


def figure(title, items, value_key, unit, path, reference=None):
    width=560;height=265;d=Drawing(width,height)
    d.add(Rect(0,0,width,height,fillColor=colors.white,strokeColor=None))
    d.add(String(12,242,title,fontName='RF-Bold',fontSize=14,fillColor=NAVY))
    left=110;right=70;chartwidth=width-left-right;maxv=max(x[value_key] for x in items)*1.08
    for i,r in enumerate(items):
        y=205-i*35; value=r[value_key]
        d.add(String(100,y+3,r['name'],fontName='RF',fontSize=10,textAnchor='end',fillColor=NAVY))
        d.add(Rect(left,y,chartwidth*value/maxv,17,fillColor=TEAL,strokeColor=None))
        label=f'{value:,.2f}' if unit=='%' else f'{value:,.0f}'
        label=label.replace(',','X').replace('.',',').replace('X','.')+unit
        d.add(String(left+chartwidth*value/maxv+7,y+3,label,fontName='RF',fontSize=10,fillColor=NAVY))
    if reference is not None:
        pos=left+chartwidth*reference/maxv
        d.add(Line(pos,48,pos,224,strokeColor=NAVY,strokeDashArray=[3,3]))
        d.add(String(12,16,'Linha tracejada: Brasil 0,9548%',fontName='RF',fontSize=10,fillColor=NAVY))
    renderPDF.drawToFile(d,str(path))
    return d


def make_figures(poppler):
    data=list(csv.DictReader((ROOT/'data/processed/territories.csv').open(encoding='utf-8')))
    rows=[{**x,'rate':float(x['rate']),'estimated_missing':float(x['estimated_missing'])} for x in data if x['year']=='2024']
    folder=ROOT/'article/figures';folder.mkdir(exist_ok=True)
    figure('Sub-registro por grande região em 2024',[x for x in rows if x['level']=='region'],'rate','%',folder/'regioes.pdf',.9548)
    # Dois painéis com unidades diferentes para evitar confundir taxa e volume.
    scratch=ROOT/'work';scratch.mkdir(exist_ok=True)
    states=[x for x in rows if x['level']=='state']
    rates=figure('Cinco maiores taxas estaduais',sorted(states,key=lambda x:x['rate'],reverse=True)[:5],'rate','%',scratch/'rates.pdf')
    volumes=figure('Cinco maiores volumes aproximados',sorted(states,key=lambda x:x['estimated_missing'],reverse=True)[:5],'estimated_missing','',scratch/'volumes.pdf')
    combined=Drawing(560,530);rates.translate(0,265);combined.add(rates);combined.add(volumes)
    renderPDF.drawToFile(combined,str(folder/'prioridades.pdf'))
    if poppler:
        for name in ['regioes','prioridades']:
            subprocess.run([poppler,'-png','-singlefile','-r','140',str(folder/(name+'.pdf')),str(folder/name)],check=True,capture_output=True)
    for name in ['regioes','prioridades']:
        if not (folder/(name+'.png')).exists():raise RuntimeError('Informe --poppler para gerar figuras PNG')


def pdf_article():
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle('BodyRF',fontName='RF',fontSize=10.4,leading=15.5,spaceAfter=9,wordWrap='LTR',splitLongWords=True))
    styles.add(ParagraphStyle('TitleRF',fontName='RF-Bold',fontSize=24,leading=29,spaceAfter=15,textColor=colors.black))
    styles.add(ParagraphStyle('H2RF',fontName='RF-Bold',fontSize=14,leading=19,spaceBefore=17,spaceAfter=9,keepWithNext=True,textColor=colors.black))
    styles.add(ParagraphStyle('H3RF',fontName='RF-Bold',fontSize=11.5,leading=16,spaceBefore=10,spaceAfter=7,keepWithNext=True,textColor=colors.black))
    styles.add(ParagraphStyle('CellRF',fontName='RF',fontSize=9,leading=12,spaceAfter=0))
    styles.add(ParagraphStyle('RefRF',fontName='RF',fontSize=8.7,leading=12.5,spaceAfter=9,splitLongWords=True))
    story=[];in_refs=False
    for b in blocks((ROOT/'article/artigo.md').read_text(encoding='utf-8')):
        if b[0]=='heading':
            if b[2]=='Referências':in_refs=True
            story.append(Paragraph(html.escape(b[2]),styles['TitleRF' if b[1]==1 else 'H2RF' if b[1]==2 else 'H3RF']))
        elif b[0]=='paragraph':story.append(Paragraph(html.escape(b[1]),styles['RefRF' if in_refs else 'BodyRF']))
        elif b[0]=='image':
            p=ROOT/'article'/b[2];img=Image(str(p));width=460;img.drawHeight=img.imageHeight*width/img.imageWidth;img.drawWidth=width
            img.keepWithNext=True
            story.append(img)
        elif b[0]=='table':
            rows=[[Paragraph(html.escape(c),styles['CellRF']) for c in r] for r in b[1]]
            n=len(rows[0]);widths=[150]+[(A4[0]-100-150)/(n-1)]*(n-1)
            t=Table(rows,colWidths=widths,repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#e1e8ef')),('GRID',(0,0),(-1,-1),.5,colors.HexColor('#d9d9d9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#f5f7fa')])]))
            story.extend([Spacer(1,6),t,Spacer(1,12)])
    def footer(c,doc):
        c.setFont('RF',8);c.setFillColor(colors.HexColor('#52677d'));c.drawString(50,28,'Registro em Foco | Pesquisa independente | Setembro de 2026');c.drawRightString(A4[0]-50,28,str(doc.page))
    SimpleDocTemplate(str(ROOT/'article/artigo.pdf'),pagesize=A4,rightMargin=50,leftMargin=50,topMargin=45,bottomMargin=48,title='Sub-registro de nascimento e prioridades territoriais no Brasil',author='Osmarsrjunior').build(story,onFirstPage=footer,onLaterPages=footer)


def docx_article():
    doc=Document();sec=doc.sections[0];sec.top_margin=Inches(.8);sec.bottom_margin=Inches(.8);sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
    sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
    for name in ['Normal','Title','Heading 1','Heading 2','Heading 3']:
        st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.2;normal.paragraph_format.space_after=Pt(8)
    doc.styles['Title'].font.size=Pt(24)
    for b in blocks((ROOT/'article/artigo.md').read_text(encoding='utf-8')):
        if b[0]=='heading':doc.add_paragraph(b[2],style='Title' if b[1]==1 else 'Heading '+str(b[1]-1))
        elif b[0]=='paragraph':doc.add_paragraph(b[1])
        elif b[0]=='image':doc.add_picture(str(ROOT/'article'/b[2]),width=Inches(6.2))
        elif b[0]=='table':
            table=doc.add_table(rows=0,cols=len(b[1][0]));table.style='Table Grid'
            for i,row in enumerate(b[1]):
                cells=table.add_row().cells
                for cell,text in zip(cells,row):
                    cell.text=text
                    for p in cell.paragraphs:
                        for r in p.runs:r.font.size=Pt(9);r.bold=(i==0)
                    tcPr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E1E8EF' if i==0 else 'FFFFFF' if i%2 else 'F5F7FA');tcPr.append(shade)
                    borders=OxmlElement('w:tcBorders')
                    for side in ['top','left','bottom','right']:
                        edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
                    tcPr.append(borders)
            doc.add_paragraph()
    doc.core_properties.author='Osmarsrjunior';doc.core_properties.title='Sub-registro de nascimento e prioridades territoriais no Brasil'
    doc.save(ROOT/'article/artigo.docx')


def html_body(text, image_prefix=''):
    result=[]
    for b in blocks(text):
        if b[0]=='heading':result.append(f'<h{b[1]}'+(' id="metodo"' if b[2]=='2 Dados e método' else '')+'>'+html.escape(b[2])+f'</h{b[1]}>')
        elif b[0]=='paragraph':
            def link_url(match):
                raw = match.group(0)
                url = raw.rstrip('.,;')
                return f'<a href="{url}">{url}</a>' + raw[len(url):]
            p=html.escape(b[1]);p=re.sub(r'https?://[^\s]+',link_url,p);result.append('<p>'+p+'</p>')
        elif b[0]=='image':result.append(f'<img src="{image_prefix}{html.escape(b[2])}" alt="{html.escape(b[1])}">')
        elif b[0]=='table':result.append('<div class="table-wrap"><table>'+''.join('<tr>'+''.join(('<th scope="col">' if i==0 else '<td>')+html.escape(c)+('</th>' if i==0 else '</td>') for c in row)+'</tr>' for i,row in enumerate(b[1]))+'</table></div>')
    return '\n'.join(result)


def make_html():
    start='<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><link rel="stylesheet" href="style.css"><link rel="icon" href="data:image/svg+xml,%3Csvg xmlns=\'http://www.w3.org/2000/svg\' viewBox=\'0 0 32 32\'%3E%3Crect width=\'32\' height=\'32\' rx=\'6\' fill=\'%23007b73\'/%3E%3C/svg%3E"><title>{title} · Registro em Foco</title></head><body><header><div class="brand">REGISTRO <strong>EM FOCO</strong></div><nav aria-label="Principal"><a href="index.html">Explorar</a><a href="artigo.html">Artigo</a><a href="dossie.html">Dossiê técnico</a></nav></header><main><article class="article">'
    end='</article></main><footer>Registro em Foco · Pesquisa independente · 2026</footer></body></html>'
    text=(ROOT/'article/artigo.md').read_text(encoding='utf-8')
    actions='<div class="doc-actions"><a href="artigo.pdf" class="button" download>Baixar PDF ↓</a><a href="artigo.md" class="button secondary" download>Texto editável ↓</a></div>'
    (ROOT/'dist/artigo.html').write_text(start.format(title='Artigo')+actions+html_body(text)+end,encoding='utf-8')
    dossie='\n\n'.join(p.read_text(encoding='utf-8') for p in sorted((ROOT/'docs').glob('0*.md')))
    (ROOT/'dist/dossie.html').write_text(start.format(title='Dossiê técnico')+html_body(dossie)+end,encoding='utf-8')
    for name in ['artigo.pdf','artigo.md']:
        shutil.copyfile(ROOT/'article'/name,ROOT/'dist'/name)
    (ROOT/'dist/figures').mkdir(exist_ok=True)
    for p in (ROOT/'article/figures').glob('*.png'):shutil.copyfile(p,ROOT/'dist/figures'/p.name)


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--poppler',default=shutil.which('pdftoppm'));parser.add_argument('--docx',action='store_true');args=parser.parse_args()
    make_figures(args.poppler);pdf_article();make_html()
    if args.docx:docx_article()
    print('Artigo PDF, texto editável e páginas HTML gerados.')


if __name__=='__main__':main()

